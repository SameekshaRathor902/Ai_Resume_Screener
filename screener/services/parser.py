"""Extract plain text from uploaded resume files (.txt, .pdf, .docx)."""
import io
import logging
import os
import shutil

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PyPdfError

log = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class ParseError(Exception):
    """Raised with a message that is safe to show to the user."""


class NoTextError(ParseError):
    """A PDF that opened fine but has no selectable text (usually a scan)."""


def ocr_available() -> bool:
    """True when pytesseract, pypdfium2 and the tesseract program are all installed."""
    try:
        import pypdfium2  # noqa: F401
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


def ocr_pdf(data: bytes, max_pages: int = 5, lang: str = "eng", page_timeout: int = 30) -> str:
    """Read a scanned PDF by rendering each page to an image and running Tesseract on it."""
    import pypdfium2 as pdfium
    import pytesseract

    try:
        pdf = pdfium.PdfDocument(data)
    except Exception as exc:  # noqa: BLE001 - pdfium raises several types
        raise ParseError("Could not open this PDF for OCR.") from exc

    pages = []
    try:
        for index in range(min(len(pdf), max_pages)):
            image = pdf[index].render(scale=2).to_pil()
            try:
                pages.append(pytesseract.image_to_string(image, lang=lang, timeout=page_timeout))
            except RuntimeError as exc:  # pytesseract raises RuntimeError on timeout
                raise ParseError("OCR took too long on this PDF.") from exc
            except pytesseract.TesseractError as exc:
                log.warning("Tesseract failed: %s", str(exc)[:200])
                raise ParseError(
                    "OCR failed. The OCR language data may not be installed on the server."
                ) from exc
    finally:
        pdf.close()

    text = "\n".join(pages).strip()
    if not text:
        raise ParseError("OCR could not find readable text in this scan.")
    return text


def extension_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def extract_text(filename: str, data: bytes) -> str:
    ext = extension_of(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ParseError("Unsupported file type. Upload a .txt, .pdf, or .docx file.")
    if not data:
        raise ParseError("The file is empty.")

    if ext == ".txt":
        text = _read_txt(data)
    elif ext == ".pdf":
        text = _read_pdf(data)
    else:
        text = _read_docx(data)

    text = text.strip()
    if not text:
        raise ParseError("No readable text found in the file.")
    return text


def _read_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001 - library raises several types
                raise ParseError("The PDF is password-protected.") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ParseError:
        raise
    except (PyPdfError, ValueError, OSError) as exc:
        raise ParseError("Could not read this PDF. It may be corrupted.") from exc

    text = "\n".join(pages)
    if not text.strip():
        raise NoTextError(
            "This PDF has no selectable text (it may be a scan). "
            "Export a text-based PDF or upload a .docx instead."
        )
    return text


def _read_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises several types
        raise ParseError("Could not read this Word file. It may be corrupted.") from exc

    lines = [p.text for p in document.paragraphs]
    # Many resumes keep skills and experience inside tables.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.extend(p.text for p in cell.paragraphs)
    return "\n".join(lines)
