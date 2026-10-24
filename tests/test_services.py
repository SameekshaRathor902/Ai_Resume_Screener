import io
import unittest

from docx import Document
from pypdf import PdfWriter

from screener.services.parser import ParseError, extract_text
from screener.services.privacy import redact_pii
from screener.services.scoring import combine_scores, cosine, score_resume, semantic_percent

JD = """
We are looking for a Python developer with experience in Flask, REST APIs, Docker and AWS.
Knowledge of machine learning and SQL is a plus. Strong communication skills.
"""


class ScoringTests(unittest.TestCase):
    def test_strong_match_scores_high(self):
        resume = "Python developer. Built REST API services in Flask, deployed with Docker on AWS. SQL, machine learning."
        result = score_resume(JD, resume)
        self.assertGreater(result["score"], 75)
        self.assertIn("Python", result["matched_skills"])
        self.assertEqual(result["missing_skills"], [])

    def test_unrelated_resume_scores_low(self):
        resume = "Experienced chef and restaurant manager. Menu design, catering, inventory and staff training."
        self.assertLess(score_resume(JD, resume)["score"], 15)

    def test_aliases_match(self):
        jd = "Need JavaScript, Kubernetes and PostgreSQL."
        resume = "I use JS, k8s and Postgres every day."
        result = score_resume(jd, resume)
        self.assertEqual(result["score"], 100.0)

    def test_word_forms_match(self):
        result = score_resume("We need someone developing software", "I am a software developer")
        self.assertIn("developing", result["matched_keywords"])

    def test_multiword_skill_is_one_term(self):
        result = score_resume("machine learning", "I like learning and I run machines")
        self.assertEqual(result["matched_keywords"], [])
        self.assertEqual(result["missing_skills"], ["Machine Learning"])

    def test_missing_skill_hurts_more_than_missing_generic_word(self):
        jd = "Python developer wanting collaboration"
        no_skill = score_resume(jd, "developer who loves collaboration")["score"]
        no_generic = score_resume(jd, "Python developer")["score"]
        self.assertGreater(no_generic, no_skill)

    def test_special_character_skills(self):
        result = score_resume("C++ and C# and .NET", "Written in c++, c#, and dotnet")
        self.assertEqual(result["score"], 100.0)

    def test_empty_input_returns_error(self):
        self.assertIn("error", score_resume("", "resume"))
        self.assertIn("error", score_resume("the and of", "resume"))

    def test_score_is_bounded(self):
        for resume in ("", "python", JD):
            score = score_resume(JD, resume or "x")["score"]
            self.assertTrue(0 <= score <= 100)

    def test_semantic_helpers(self):
        self.assertAlmostEqual(cosine([1, 0], [1, 0]), 1.0)
        self.assertEqual(cosine([0, 0], [1, 0]), 0.0)
        self.assertEqual(semantic_percent(0.95), 100.0)
        self.assertEqual(semantic_percent(0.1), 0.0)
        self.assertEqual(combine_scores(80.0, None), 80.0)
        self.assertEqual(combine_scores(100.0, 0.0), 70.0)


class PrivacyTests(unittest.TestCase):
    def test_removes_email_phone_and_links(self):
        text = "Jane jane.doe@mail.com +91 98765 43210 https://linkedin.com/in/jane github.com/jane"
        out = redact_pii(text)
        self.assertNotIn("jane.doe", out)
        self.assertNotIn("98765", out)
        self.assertNotIn("linkedin.com", out)
        self.assertNotIn("github.com", out)

    def test_keeps_date_ranges(self):
        self.assertIn("2019 - 2023", redact_pii("Engineer 2019 - 2023"))


class ParserTests(unittest.TestCase):
    def test_txt(self):
        self.assertEqual(extract_text("cv.TXT", b"hello world"), "hello world")

    def test_unsupported_type(self):
        with self.assertRaises(ParseError):
            extract_text("cv.exe", b"x")

    def test_empty_file(self):
        with self.assertRaises(ParseError):
            extract_text("cv.txt", b"")

    def test_docx_reads_paragraphs_and_tables(self):
        doc = Document()
        doc.add_paragraph("Paragraph skill: Python")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Table skill: Docker"
        buffer = io.BytesIO()
        doc.save(buffer)
        text = extract_text("cv.docx", buffer.getvalue())
        self.assertIn("Python", text)
        self.assertIn("Docker", text)

    def test_blank_pdf_reports_no_text(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        with self.assertRaises(ParseError) as ctx:
            extract_text("scan.pdf", buffer.getvalue())
        self.assertIn("no selectable text", str(ctx.exception))

    def test_corrupt_pdf(self):
        with self.assertRaises(ParseError):
            extract_text("bad.pdf", b"not a pdf at all")


if __name__ == "__main__":
    unittest.main()
